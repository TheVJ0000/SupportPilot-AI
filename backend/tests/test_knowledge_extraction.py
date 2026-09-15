from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document

from app.knowledge import extraction
from app.knowledge.errors import ExtractionError
from app.knowledge.extraction import extract_docx, extract_pdf, extract_text


def docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Opening paragraph")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Plan"
    table.cell(0, 1).text = "Free"
    document.add_paragraph("Closing paragraph")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_text_extraction_preserves_line_locators_and_handles_utf8_bom() -> None:
    segments = extract_text(
        b"\xef\xbb\xbfFirst line\r\nSecond line\r\n\r\nFourth line", kind="text"
    )

    assert [segment.text for segment in segments] == ["First line\nSecond line", "Fourth line"]
    assert segments[0].locator == {"kind": "text", "line_start": 1, "line_end": 2}
    assert segments[1].locator == {"kind": "text", "line_start": 4, "line_end": 4}


@pytest.mark.parametrize(
    ("content", "error_code"),
    [
        (b"\xff\xfeinvalid", "unsupported_encoding"),
        (b"text\x00binary", "invalid_file_content"),
        (b" \n\t\n", "no_extractable_text"),
    ],
)
def test_text_extraction_rejects_unsafe_or_empty_content(content: bytes, error_code: str) -> None:
    with pytest.raises(ExtractionError, match=error_code):
        extract_text(content, kind="markdown")


def test_docx_extracts_paragraphs_and_tables_in_document_order() -> None:
    segments = extract_docx(docx_bytes())

    assert [segment.text for segment in segments] == [
        "Opening paragraph",
        "Plan\tFree",
        "Closing paragraph",
    ]
    assert [segment.locator["block_start"] for segment in segments] == [1, 2, 3]


def test_docx_rejects_malformed_container() -> None:
    with pytest.raises(ExtractionError, match="malformed_document"):
        extract_docx(b"not-a-zip")


def test_docx_rejects_traversal_entry() -> None:
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<document/>")
        archive.writestr("../outside.xml", "<outside/>")

    with pytest.raises(ExtractionError, match="malformed_document"):
        extract_docx(output.getvalue())


def test_docx_rejects_external_relationships() -> None:
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<document/>")
        archive.writestr(
            "word/_rels/document.xml.rels",
            '<Relationship TargetMode="External" Target="https://example.test"/>',
        )

    with pytest.raises(ExtractionError, match="invalid_file_content"):
        extract_docx(output.getvalue())


def test_docx_rejects_excessive_entries(monkeypatch) -> None:
    monkeypatch.setattr(extraction, "MAX_DOCX_ENTRIES", 2)
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<document/>")
        archive.writestr("word/styles.xml", "<styles/>")

    with pytest.raises(ExtractionError, match="document_too_large"):
        extract_docx(output.getvalue())


def test_docx_rejects_oversized_member(monkeypatch) -> None:
    monkeypatch.setattr(extraction, "MAX_DOCX_MEMBER_BYTES", 8)
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<document/>too-large")

    with pytest.raises(ExtractionError, match="document_too_large"):
        extract_docx(output.getvalue())


def test_docx_rejects_excessive_total_expansion(monkeypatch) -> None:
    monkeypatch.setattr(extraction, "MAX_DOCX_UNCOMPRESSED_BYTES", 24)
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<document/>")
        archive.writestr("word/styles.xml", "<styles/>")

    with pytest.raises(ExtractionError, match="document_too_large"):
        extract_docx(output.getvalue())


class FakePdfPage:
    def __init__(self, text: str | None) -> None:
        self._text = text

    def extract_text(self) -> str | None:
        return self._text


class FakePdfReader:
    def __init__(self, pages: list[FakePdfPage], encrypted: bool = False) -> None:
        self.pages = pages
        self.is_encrypted = encrypted


def test_pdf_rejects_invalid_signature() -> None:
    with pytest.raises(ExtractionError, match="invalid_file_content"):
        extract_pdf(b"not a pdf")


def test_pdf_rejects_malformed_signed_content() -> None:
    with pytest.raises(ExtractionError, match="malformed_document"):
        extract_pdf(b"%PDF-not-a-real-document")


def test_pdf_rejects_encryption(monkeypatch) -> None:
    monkeypatch.setattr(extraction, "PdfReader", lambda *_args, **_kwargs: FakePdfReader([], True))
    with pytest.raises(ExtractionError, match="encrypted_pdf"):
        extract_pdf(b"%PDF-synthetic")


def test_pdf_rejects_excessive_pages(monkeypatch) -> None:
    monkeypatch.setattr(extraction, "MAX_PDF_PAGES", 1)
    monkeypatch.setattr(
        extraction,
        "PdfReader",
        lambda *_args, **_kwargs: FakePdfReader([FakePdfPage("one"), FakePdfPage("two")]),
    )
    with pytest.raises(ExtractionError, match="too_many_pages"):
        extract_pdf(b"%PDF-synthetic")


def test_pdf_rejects_excessive_extracted_text(monkeypatch) -> None:
    monkeypatch.setattr(extraction, "MAX_EXTRACTED_TEXT_CHARS", 5)
    monkeypatch.setattr(
        extraction,
        "PdfReader",
        lambda *_args, **_kwargs: FakePdfReader([FakePdfPage("too much text")]),
    )
    with pytest.raises(ExtractionError, match="document_too_large"):
        extract_pdf(b"%PDF-synthetic")


def test_pdf_preserves_one_based_page_locators(monkeypatch) -> None:
    monkeypatch.setattr(
        extraction,
        "PdfReader",
        lambda *_args, **_kwargs: FakePdfReader([FakePdfPage("one"), FakePdfPage("two")]),
    )

    segments = extract_pdf(b"%PDF-synthetic")

    assert segments[1].locator == {"kind": "pdf", "page_start": 2, "page_end": 2}


def test_pdf_rejects_image_only_or_empty_text(monkeypatch) -> None:
    monkeypatch.setattr(
        extraction, "PdfReader", lambda *_args, **_kwargs: FakePdfReader([FakePdfPage(None)])
    )
    with pytest.raises(ExtractionError, match="no_extractable_text"):
        extract_pdf(b"%PDF-synthetic")
